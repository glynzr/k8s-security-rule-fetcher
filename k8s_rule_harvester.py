#!/usr/bin/env python3
"""
k8s_rule_harvester.py — collect Kubernetes security check rules from open-source
scanners, normalise them into one schema, and export a review workbook so a team
can pick the rules that become its internal K8s security guideline.

Sources (pulled straight from upstream repos, pinned to the fetched commit):
  kube-bench        CIS Kubernetes Benchmark (generic, EKS/AKS/GKE, OpenShift, RKE2, k3s)
  kubescape         Kubescape regolibrary controls (+ NSA, MITRE, CIS, SOC2 framework mapping)
  trivy             Trivy / Aqua trivy-checks (KSV-*, incl. CIS ids per distro)
  checkov           Checkov Kubernetes checks (CKV_K8S_*, CKV2_K8S_*)
  kube-linter       StackRox kube-linter built-in checks
  polaris           Fairwinds Polaris checks (+ default severity)
  kyverno           Kyverno policy library (validating / image-verification policies)
  prowler           Prowler Kubernetes provider checks
  gatekeeper        OPA Gatekeeper library constraint templates

Subcommands
  harvest   fetch + normalise + export  (default)
  draft     read the reviewed workbook and generate a draft company guideline (Markdown)

Typical flow
  python3 k8s_rule_harvester.py harvest --k8s-version 1.30 -o out/
  # team reviews out/k8s-security-rules.xlsx: fills Decision / Guideline ID / Owner / Notes
  python3 k8s_rule_harvester.py draft out/k8s-security-rules.xlsx -o out/k8s-security-guideline-draft.md

Requirements: Python 3.9+, git, PyYAML, openpyxl  (pip install pyyaml openpyxl)
"""
from __future__ import annotations

import argparse
import ast
import csv
import datetime as dt
import io
import json
import logging
import re
import shutil
import subprocess
import sys
import tarfile
import urllib.request
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Callable, Iterable

try:
    import yaml
except ImportError:  # pragma: no cover
    sys.exit("PyYAML is required: pip install pyyaml")

log = logging.getLogger("harvester")

# --------------------------------------------------------------------------- #
# Source registry
# --------------------------------------------------------------------------- #
@dataclass
class SourceSpec:
    key: str
    repo: str                       # owner/name on GitHub
    sparse: list[str] | None        # sparse-checkout paths (None = whole repo)
    description: str
    parser: str                     # name of parser function


SOURCES: dict[str, SourceSpec] = {
    s.key: s for s in [
        SourceSpec("kube-bench", "aquasecurity/kube-bench", ["cfg"],
                   "CIS Kubernetes Benchmark checks (node/control-plane)", "parse_kube_bench"),
        SourceSpec("kubescape", "kubescape/regolibrary", ["controls", "frameworks"],
                   "Kubescape controls with NSA/MITRE/CIS/SOC2 mapping", "parse_kubescape"),
        SourceSpec("trivy", "aquasecurity/trivy-checks", ["checks/kubernetes"],
                   "Trivy misconfiguration checks (KSV)", "parse_trivy"),
        SourceSpec("checkov", "bridgecrewio/checkov", ["checkov/kubernetes/checks"],
                   "Checkov Kubernetes manifest checks", "parse_checkov"),
        SourceSpec("kube-linter", "stackrox/kube-linter", ["pkg/builtinchecks/yamls"],
                   "kube-linter built-in checks", "parse_kube_linter"),
        SourceSpec("polaris", "FairwindsOps/polaris", ["pkg/config"],
                   "Polaris workload best-practice checks", "parse_polaris"),
        SourceSpec("kyverno", "kyverno/policies", None,
                   "Kyverno policy library (validating policies)", "parse_kyverno"),
        SourceSpec("prowler", "prowler-cloud/prowler", ["prowler/providers/kubernetes/services"],
                   "Prowler Kubernetes provider checks", "parse_prowler"),
        SourceSpec("gatekeeper", "open-policy-agent/gatekeeper-library", ["library"],
                   "OPA Gatekeeper constraint template library", "parse_gatekeeper"),
    ]
}

# --------------------------------------------------------------------------- #
# Normalised rule model
# --------------------------------------------------------------------------- #
REVIEW_COLUMNS = ["decision", "guideline_id", "guideline_text", "owner", "review_notes"]
DECISIONS = ["Adopt", "Adapt", "Reject", "Not applicable", "Needs discussion"]
SEVERITY_ORDER = ["critical", "high", "medium", "low", "info", "unknown"]


@dataclass
class Rule:
    source: str
    rule_id: str
    title: str
    description: str = ""
    rationale: str = ""
    remediation: str = ""
    severity: str = "unknown"          # normalised
    raw_severity: str = ""
    check_type: str = "automated"      # automated | manual
    scope: str = ""                    # cluster-config (node/control plane) | manifest | rbac | runtime-api
    resource_kinds: str = ""
    benchmark: str = ""                # e.g. cis-1.12, eks-1.8.0, kyverno category
    cis_refs: str = ""
    frameworks: str = ""
    audit: str = ""                    # how the check is performed (kube-bench audit cmd etc.)
    references: str = ""
    source_url: str = ""
    deprecated_note: str = ""
    domain: str = ""
    topics: str = ""
    # review columns (filled by humans)
    decision: str = ""
    guideline_id: str = ""
    guideline_text: str = ""
    owner: str = ""
    review_notes: str = ""

    @property
    def uid(self) -> str:
        b = f"{self.benchmark}:" if self.source == "kube-bench" else ""
        return f"{self.source}:{b}{self.rule_id}"


# --------------------------------------------------------------------------- #
# Fetching
# --------------------------------------------------------------------------- #
@dataclass
class Checkout:
    path: Path
    repo: str
    commit: str

    def url(self, file: Path) -> str:
        rel = file.relative_to(self.path).as_posix()
        return f"https://github.com/{self.repo}/blob/{self.commit}/{rel}"


def _run(cmd: list[str], cwd: Path | None = None) -> str:
    return subprocess.run(cmd, cwd=cwd, check=True, capture_output=True, text=True).stdout.strip()


def fetch(spec: SourceSpec, cache: Path, ref: str | None, refresh: bool, offline: bool) -> Checkout:
    dest = cache / spec.key
    if dest.exists() and (refresh and not offline):
        shutil.rmtree(dest)
    if not dest.exists():
        if offline:
            raise RuntimeError(f"{spec.key}: not in cache ({dest}) and --offline given")
        if shutil.which("git"):
            _git_clone(spec, dest, ref)
        else:
            _tarball(spec, dest, ref)
    commit = "unknown"
    if (dest / ".git").exists():
        commit = _run(["git", "rev-parse", "HEAD"], dest)
    elif (dest / ".commit").exists():
        commit = (dest / ".commit").read_text().strip()
    return Checkout(dest, spec.repo, commit)


def _git_clone(spec: SourceSpec, dest: Path, ref: str | None) -> None:
    url = f"https://github.com/{spec.repo}.git"
    cmd = ["git", "clone", "-q", "--depth", "1", "--filter=blob:none"]
    if spec.sparse:
        cmd.append("--sparse")
    if ref:
        cmd += ["--branch", ref]
    log.info("cloning %s%s", spec.repo, f"@{ref}" if ref else "")
    _run(cmd + [url, str(dest)])
    if spec.sparse:
        _run(["git", "sparse-checkout", "set", *spec.sparse], dest)


def _tarball(spec: SourceSpec, dest: Path, ref: str | None) -> None:
    """Fallback when git is unavailable: download a codeload tarball."""
    url = f"https://codeload.github.com/{spec.repo}/tar.gz/{ref or 'HEAD'}"
    log.info("downloading %s", url)
    with urllib.request.urlopen(url, timeout=300) as r:
        data = r.read()
    dest.mkdir(parents=True)
    with tarfile.open(fileobj=io.BytesIO(data), mode="r:gz") as tf:
        root = tf.getmembers()[0].name.split("/")[0]
        for m in tf.getmembers():
            if not m.isfile():
                continue
            rel = m.name[len(root) + 1:]
            if spec.sparse and not any(rel.startswith(p.rstrip("/") + "/") for p in spec.sparse):
                continue
            target = dest / rel
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(tf.extractfile(m).read())
    (dest / ".commit").write_text(ref or "HEAD (tarball)")


# --------------------------------------------------------------------------- #
# Helpers
# --------------------------------------------------------------------------- #
def clean(text) -> str:
    if text is None:
        return ""
    if isinstance(text, (list, tuple)):
        text = "\n".join(str(t) for t in text)
    text = str(text).replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def norm_severity(raw) -> str:
    if raw is None or raw == "":
        return "unknown"
    s = str(raw).strip().lower()
    try:  # numeric scores (Kubescape baseScore 0-10)
        v = float(s)
        return "critical" if v >= 9 else "high" if v >= 7 else "medium" if v >= 4 else "low" if v > 0 else "info"
    except ValueError:
        pass
    return {
        "critical": "critical", "high": "high", "danger": "high", "error": "high",
        "medium": "medium", "moderate": "medium", "warning": "medium",
        "low": "low", "info": "info", "informational": "info", "ignore": "info",
    }.get(s, "unknown")


def load_yaml_docs(path: Path) -> list:
    try:
        return [d for d in yaml.safe_load_all(path.read_text(encoding="utf-8", errors="replace")) if d]
    except yaml.YAMLError as e:
        log.debug("yaml error in %s: %s", path, e)
        return []


COMPONENTS = [  # (label, regex over file name / source)
    ("API server", r"apiserver|kube-apiserver"),
    ("controller manager", r"controller.?manager"),
    ("scheduler", r"scheduler"),
    ("etcd", r"\betcd"),
    ("kubelet", r"kubelet"),
    ("kube-proxy", r"kube.?proxy"),
]


def detect_component(text: str) -> str:
    return next((c for c, rx in COMPONENTS if re.search(rx, text, re.I)), "")


CIS_RE = re.compile(r"\bCIS[\w .-]*?\b(\d+\.\d+(?:\.\d+){1,2})\b", re.I)


def cis_from_text(text: str) -> list[str]:
    return sorted({m.group(1) for m in CIS_RE.finditer(text or "")})


# --------------------------------------------------------------------------- #
# Parsers — each yields Rule objects
# --------------------------------------------------------------------------- #
def _kb_select_benchmarks(cfg: Path, wanted: list[str], k8s_version: str | None) -> list[str]:
    available = sorted(p.name for p in cfg.iterdir() if p.is_dir())
    if wanted == ["all"]:
        return available
    chosen: list[str] = []
    if k8s_version:
        conf = yaml.safe_load((cfg / "config.yaml").read_text()) or {}
        mapping = conf.get("version_mapping", {}) or {}
        if k8s_version in mapping:
            chosen.append(mapping[k8s_version])
        else:
            log.warning("kube-bench: no version_mapping for %s; falling back to latest CIS", k8s_version)
    for w in wanted:
        if w == "latest":
            generic = [b for b in available if re.fullmatch(r"cis-\d+\.\d+", b)]
            generic.sort(key=lambda b: tuple(int(x) for x in b[4:].split(".")))
            if generic and not chosen:
                chosen.append(generic[-1])
        elif w in available:
            chosen.append(w)
        else:
            log.warning("kube-bench: benchmark '%s' not found. Available: %s", w, ", ".join(available))
    return list(dict.fromkeys(chosen))


KB_SCOPE = {"master": "control plane node", "node": "worker node", "etcd": "etcd",
            "controlplane": "control plane (authn/logging)", "policies": "cluster policies",
            "managedservices": "managed service"}


def parse_kube_bench(co: Checkout, opts) -> Iterable[Rule]:
    cfg = co.path / "cfg"
    benches = _kb_select_benchmarks(cfg, opts.kube_bench_benchmarks, opts.k8s_version)
    log.info("kube-bench benchmarks: %s", ", ".join(benches))
    for bench in benches:
        for f in sorted((cfg / bench).glob("*.yaml")):
            if f.name == "config.yaml":
                continue
            for doc in load_yaml_docs(f):
                ftype = str(doc.get("type", f.stem))
                for g in doc.get("groups", []) or []:
                    gtext = clean(g.get("text"))
                    for c in g.get("checks", []) or []:
                        text = clean(c.get("text"))
                        manual = "(manual)" in text.lower() or str(c.get("type", "")).lower() == "manual"
                        title = re.sub(r"\s*\((Automated|Manual|Scored|Not Scored)\)\s*$", "", text, flags=re.I)
                        yield Rule(
                            source="kube-bench", rule_id=str(c.get("id")), title=title,
                            description=f"[{gtext}] {title}",
                            remediation=clean(c.get("remediation")),
                            severity="unknown",
                            raw_severity="scored" if c.get("scored") else "not scored",
                            check_type="manual" if manual else "automated",
                            scope=f"cluster-config: {KB_SCOPE.get(ftype, ftype)}",
                            benchmark=bench, cis_refs=f"{bench}:{c.get('id')}",
                            frameworks=f"CIS ({bench})", audit=clean(c.get("audit")),
                            source_url=co.url(f), topics=gtext,  # topics overwritten later; gtext kept in description
                        )


def parse_kubescape(co: Checkout, opts) -> Iterable[Rule]:
    fw_map: dict[str, list[str]] = defaultdict(list)
    cis_map: dict[str, list[str]] = defaultdict(list)
    for f in sorted((co.path / "frameworks").glob("*.json")):
        try:
            fw = json.loads(f.read_text())
        except json.JSONDecodeError:
            continue
        name = fw.get("name", f.stem)
        if name.lower() in {"allcontrols", "__yamlscan", "clusterscan", "workloadscan"}:
            continue
        for ac in fw.get("activeControls", []) or []:
            cid = ac.get("controlID")
            if not cid:
                continue
            fw_map[cid].append(name)
            pname = (ac.get("patch") or {}).get("name", "")
            m = re.match(r"CIS-(\d+(?:\.\d+)+)", pname)
            if m:
                cis_map[cid].append(f"{name}:{m.group(1)}")
    for f in sorted((co.path / "controls").glob("*.json")):
        try:
            c = json.loads(f.read_text())
        except json.JSONDecodeError:
            continue
        cid = c.get("controlID") or f.stem
        attrs = c.get("attributes", {}) or {}
        cat = c.get("category", {}) or {}
        scope = ", ".join((c.get("scanningScope") or {}).get("matches", []))
        mitre = attrs.get("microsoftMitreColumns", []) or []
        fws = sorted(set(fw_map.get(cid, [])))
        if mitre:
            fws.append("MITRE: " + ", ".join(mitre))
        yield Rule(
            source="kubescape", rule_id=cid, title=clean(c.get("name")),
            description=clean(c.get("description")),
            rationale=ld if len(ld := clean(c.get("long_description"))) > 80 else "",
            remediation=clean(c.get("remediation")),
            severity=norm_severity(c.get("baseScore")), raw_severity=f"baseScore {c.get('baseScore', '')}",
            check_type="manual" if "manual" in str(attrs.get("actionRequired", "")).lower() else "automated",
            scope=f"kubescape scope: {scope}" if scope else "",
            benchmark=clean(cat.get("name")) + (f" / {cat['subCategory']['name']}" if isinstance(cat.get("subCategory"), dict) else ""),
            cis_refs="; ".join(sorted(set(cis_map.get(cid, [])))),
            frameworks="; ".join(fws), audit=clean(c.get("test")),
            references=f"https://hub.armosec.io/docs/{cid.lower()}",
            source_url=co.url(f),
        )


def _rego_metadata(text: str) -> dict:
    lines = text.splitlines()
    try:
        start = next(i for i, l in enumerate(lines) if l.strip() == "# METADATA")
    except StopIteration:
        return {}
    block = []
    for l in lines[start + 1:]:
        if not l.startswith("#"):
            break
        block.append(l[2:] if l.startswith("# ") else l[1:])
    try:
        return yaml.safe_load("\n".join(block)) or {}
    except yaml.YAMLError:
        return {}


def parse_trivy(co: Checkout, opts) -> Iterable[Rule]:
    for f in sorted((co.path / "checks/kubernetes").rglob("*.rego")):
        if f.name.endswith("_test.rego"):
            continue
        md = _rego_metadata(f.read_text(encoding="utf-8", errors="replace"))
        if not md or not md.get("custom"):
            continue
        cu = md["custom"]
        if cu.get("deprecated"):
            continue
        frameworks = cu.get("frameworks") or {}
        cis = [f"{k}:{i}" for k, ids in frameworks.items() if "cis" in k.lower() for i in (ids or [])]
        kinds = sorted({s.get("kind") for sel in (cu.get("input") or {}).get("selector", []) or []
                        for s in (sel.get("subtypes") or []) if s.get("kind")})
        comp = detect_component(f.stem) if (not kinds or kinds == ["nodeinfo"]) else ""
        cfg_check = bool(comp) or kinds == ["nodeinfo"]
        desc = clean(md.get("description"))
        yield Rule(
            source="trivy", rule_id=str(cu.get("id") or cu.get("avd_id")), title=clean(md.get("title")),
            description=f"[{comp}] {desc}" if comp else desc,
            remediation=clean(cu.get("recommended_actions") or cu.get("recommended_action")),
            severity=norm_severity(cu.get("severity")), raw_severity=str(cu.get("severity", "")),
            scope=f"cluster-config: {comp or 'node'}" if cfg_check else "manifest",
            resource_kinds=", ".join(kinds), benchmark="trivy-checks",
            cis_refs="; ".join(cis),
            frameworks="; ".join(k for k in frameworks if k != "default"),
            references="\n".join(md.get("related_resources") or []),
            source_url=co.url(f),
        )


def _py_const(node):
    try:
        return ast.literal_eval(node)
    except Exception:
        return None


def parse_checkov(co: Checkout, opts) -> Iterable[Rule]:
    base = co.path / "checkov/kubernetes/checks"
    for f in sorted((base / "resource/k8s").glob("*.py")):
        if f.name.startswith("__"):
            continue
        src = f.read_text(encoding="utf-8", errors="replace")
        try:
            tree = ast.parse(src)
        except SyntaxError:
            continue
        found: dict[str, object] = {}
        for node in ast.walk(tree):
            if isinstance(node, ast.Assign) and len(node.targets) == 1 and isinstance(node.targets[0], ast.Name):
                n = node.targets[0].id
                if n in ("name", "id", "supported_kind", "supported_entities", "supported_resources"):
                    v = _py_const(node.value)
                    if v is not None:
                        found.setdefault(n, v)
            if isinstance(node, ast.Call) and getattr(node.func, "attr", "") == "__init__":
                for kw in node.keywords:
                    if kw.arg in ("name", "id", "supported_entities") and kw.arg not in found:
                        v = _py_const(kw.value)
                        if v is not None:
                            found[kw.arg] = v
        cid, name = found.get("id"), found.get("name")
        if not cid or not name:
            continue
        kinds = found.get("supported_kind") or found.get("supported_entities") or found.get("supported_resources") or ()
        kinds = [kinds] if isinstance(kinds, str) else list(kinds)
        comments = "\n".join(l.strip()[1:].strip() for l in src.splitlines() if l.strip().startswith("#"))
        refs = "\n".join(sorted(set(re.findall(r"https?://\S+", comments))))
        cis = re.findall(r"CIS[- ]?(\d+\.\d+)\s+(\d+(?:\.\d+)+)", comments)
        comp = ""
        if re.search(r"--[a-z-]+", str(name)):  # CLI-flag check -> which binary does it inspect?
            m = re.search(r'["\'](kube-apiserver|kube-controller-manager|kube-scheduler|etcd|kubelet)["\']', src)
            comp = detect_component(m.group(1) if m else f.stem)
        dep = ""
        if any("PodSecurityPolicy" in str(k) for k in kinds) or "PSP" in f.stem:
            dep = "PodSecurityPolicy was removed in Kubernetes 1.25 — map to Pod Security Admission / policy engine instead."
        yield Rule(
            source="checkov", rule_id=str(cid), title=clean(name),
            description=f"[{comp}] {clean(name)}" if comp else clean(name), severity="unknown",
            raw_severity="(severity only in commercial Prisma Cloud)",
            scope=f"cluster-config: {comp}" if comp else "rbac" if any(k in ("Role", "ClusterRole", "RoleBinding", "ClusterRoleBinding") for k in kinds) else "manifest",
            resource_kinds=", ".join(map(str, kinds)), benchmark="checkov",
            cis_refs="; ".join(f"cis-{v}:{i}" for v, i in cis),
            references=refs + f"\nhttps://docs.prismacloud.io/en/enterprise-edition/policy-reference/kubernetes-policies",
            source_url=co.url(f), deprecated_note=dep,
        )
    for f in sorted((base / "graph_checks").glob("*.yaml")):
        for doc in load_yaml_docs(f):
            md = doc.get("metadata") or {}
            if not md.get("id"):
                continue
            yield Rule(
                source="checkov", rule_id=str(md["id"]), title=clean(md.get("name")),
                description=clean(md.get("name")), severity="unknown",
                raw_severity="(severity only in commercial Prisma Cloud)",
                scope="manifest (graph/relationship check)", benchmark="checkov",
                source_url=co.url(f),
            )


def parse_kube_linter(co: Checkout, opts) -> Iterable[Rule]:
    for f in sorted((co.path / "pkg/builtinchecks/yamls").glob("*.yaml")):
        for d in load_yaml_docs(f):
            desc = clean(d.get("description"))
            kinds = (d.get("scope") or {}).get("objectKinds") or []
            yield Rule(
                source="kube-linter", rule_id=str(d.get("name")), title=str(d.get("name")).replace("-", " ").capitalize(),
                description=desc, remediation=clean(d.get("remediation")),
                scope="manifest", resource_kinds=", ".join(kinds), benchmark="kube-linter",
                cis_refs="; ".join(f"cis:{c}" for c in cis_from_text(desc)),
                audit=f"template: {d.get('template')}" + (f" params: {json.dumps(d.get('params'))}" if d.get("params") else ""),
                references="https://docs.kubelinter.io/#/generated/checks",
                source_url=co.url(f),
            )


def parse_polaris(co: Checkout, opts) -> Iterable[Rule]:
    base = co.path / "pkg/config"
    defaults = (yaml.safe_load((base / "default.yaml").read_text()) or {}).get("checks", {}) if (base / "default.yaml").exists() else {}
    for f in sorted((base / "checks").glob("*.yaml")):
        docs = load_yaml_docs(f)
        if not docs:
            continue
        d = docs[0]
        name = f.stem
        raw = defaults.get(name, "not in default config")
        yield Rule(
            source="polaris", rule_id=name,
            title=clean(d.get("failureMessage") or d.get("FailureMessage") or name),
            description=f"Pass: {clean(d.get('successMessage'))} | Fail: {clean(d.get('failureMessage') or d.get('FailureMessage'))}",
            severity=norm_severity(raw), raw_severity=str(raw),
            scope="manifest", resource_kinds=clean(d.get("target")),
            benchmark=f"polaris/{clean(d.get('category'))}",
            references=f"https://polaris.docs.fairwinds.com/checks/{clean(d.get('category')).lower()}/",
            source_url=co.url(f),
        )


KYVERNO_KEEP_KINDS = ("ClusterPolicy", "Policy", "ValidatingPolicy", "NamespacedValidatingPolicy",
                      "ImageValidatingPolicy", "NamespacedImageValidatingPolicy")


def parse_kyverno(co: Checkout, opts) -> Iterable[Rule]:
    seen: set[str] = set()
    for f in sorted(co.path.rglob("*.yaml")):
        rel = f.relative_to(co.path).as_posix()
        if any(part.startswith(".") for part in rel.split("/")) or f.name in ("kustomization.yaml", "artifacthub-pkg.yml"):
            continue
        for d in load_yaml_docs(f):
            if not isinstance(d, dict) or d.get("kind") not in KYVERNO_KEEP_KINDS:
                continue
            md = d.get("metadata") or {}
            ann = md.get("annotations") or {}
            title = ann.get("policies.kyverno.io/title")
            if not title:
                continue
            if d.get("kind") in ("ClusterPolicy", "Policy"):
                rules = (d.get("spec") or {}).get("rules") or []
                if rules and not any("validate" in r or "verifyImages" in r for r in rules):
                    continue  # mutate/generate-only
            name = md.get("name", f.stem)
            title = re.sub(r"\s+in\s+(Validating|ImageValidating|Mutating|Generating)Policy$", "", clean(title))
            category = re.sub(r"\s+in\s+\w+Policy$", "", clean(ann.get("policies.kyverno.io/category")))
            key = name
            if key in seen:
                continue
            seen.add(key)
            yield Rule(
                source="kyverno", rule_id=name, title=title,
                description=clean(ann.get("policies.kyverno.io/description")),
                severity=norm_severity(ann.get("policies.kyverno.io/severity")),
                raw_severity=clean(ann.get("policies.kyverno.io/severity")),
                scope="admission (manifest)", resource_kinds=clean(ann.get("policies.kyverno.io/subject")),
                benchmark=f"kyverno/{category or rel.split('/')[0]}",
                frameworks="Pod Security Standards" if "Pod Security" in category else "",
                audit=f"{d.get('kind')} (min Kyverno {ann.get('policies.kyverno.io/minversion', '?')})",
                references=f"https://kyverno.io/policies/?policytypes={name}",
                source_url=co.url(f),
            )


def parse_prowler(co: Checkout, opts) -> Iterable[Rule]:
    for f in sorted((co.path / "prowler/providers/kubernetes/services").rglob("*.metadata.json")):
        try:
            m = json.loads(f.read_text())
        except json.JSONDecodeError:
            continue
        rem = m.get("Remediation") or {}
        rec = rem.get("Recommendation") or {}
        code = rem.get("Code") or {}
        urls = [u for u in [m.get("RelatedUrl"), rec.get("Url"), *(m.get("AdditionalURLs") or [])] if u]
        rtext = clean(rec.get("Text"))
        if code.get("Other"):
            rtext += "\n\nSteps:\n" + clean(code["Other"])
        yield Rule(
            source="prowler", rule_id=m.get("CheckID", f.stem), title=clean(m.get("CheckTitle")),
            description=clean(m.get("Description")), rationale=clean(m.get("Risk")),
            remediation=rtext.strip(),
            severity=norm_severity(m.get("Severity")), raw_severity=str(m.get("Severity", "")),
            scope=f"runtime-api: {m.get('ServiceName', '')}", resource_kinds=clean(m.get("ResourceType")),
            benchmark=f"prowler/{m.get('ServiceName', '')}",
            frameworks="; ".join(m.get("Categories") or []),
            audit=clean(m.get("Notes")),
            references="\n".join(urls), source_url=co.url(f),
        )


def parse_gatekeeper(co: Checkout, opts) -> Iterable[Rule]:
    for f in sorted((co.path / "library").rglob("template.yaml")):
        for d in load_yaml_docs(f):
            if d.get("kind") != "ConstraintTemplate":
                continue
            md = d.get("metadata") or {}
            ann = md.get("annotations") or {}
            rel = f.relative_to(co.path / "library").parts
            bundle = ann.get("metadata.gatekeeper.sh/bundle", "")
            yield Rule(
                source="gatekeeper", rule_id=md.get("name", "/".join(rel[:-1])),
                title=clean(ann.get("metadata.gatekeeper.sh/title") or md.get("name")),
                description=clean(ann.get("description")),
                scope="admission (manifest)", benchmark=f"gatekeeper/{rel[0]}",
                frameworks=f"bundle: {bundle}" if bundle else "",
                references="https://open-policy-agent.github.io/gatekeeper-library/website/",
                source_url=co.url(f),
            )


PARSERS: dict[str, Callable] = {n: globals()[n] for n in {s.parser for s in SOURCES.values()}}

# --------------------------------------------------------------------------- #
# Classification: domain + topics (to group equivalent rules across tools)
# --------------------------------------------------------------------------- #
DOMAINS: list[tuple[str, str]] = [  # first match wins; order matters
    ("Host files: permissions & ownership", r"\b(file )?permissions?\b.*\b(set to|restrictive)|ownership is set|\bchmod\b|\bchown\b"),
    ("Known vulnerabilities & legacy components", r"cve-\d{4}|outdated|tiller|helm v2|dashboard"),
    ("Worker nodes: kubelet & kube-proxy", r"\[kubelet\]|\[kube-proxy\]|kubelet|kube-proxy|worker node|streaming-connection|hostname-override|protect-kernel|make-iptables|event-qps|rotate-certificates|read-only-port"),
    ("etcd", r"\[etcd\]|\betcd\b|peer-cert|peer-client|client-cert-auth|--cert-file and --key-file"),
    ("Control plane: controller manager & scheduler", r"\[controller manager\]|\[scheduler\]|controller.?manager|kube-scheduler|terminated-pod-gc|service-account-private-key|root-ca-file|use-service-account-credentials"),
    ("Audit logging & monitoring", r"\baudit\b|logging|\blogs?\b|delete .*events"),
    ("Control plane: API server", r"\[api server\]|api.?server|admission (plugin|control)|--anonymous-auth|insecure.port|secure-port|service-account-(key|lookup)|basic-auth|token-auth|--profiling|client-ca-file"),
    ("Managed Kubernetes (EKS/AKS/GKE) & cloud", r"\b(eks|aks|gke|iam|aws|azure|gcp|google cloud|workload identity|metadata api|imds|karpenter)\b|private (endpoint|nodes)|public access"),
    ("Secrets & encryption", r"secret|credential|encrypt|encyption|\bkms\b|sensitive|password"),
    ("RBAC & identity", r"\brbac\b|clusterrole|rolebinding|\brole\b|cluster-admin|impersonat|service ?account|authoriz|authenticat|\bverbs?\b|escalate|\bbind\b|kubeconfig|system:(masters|anonymous|authenticated)"),
    ("Network security", r"network ?polic|hostport|host port|ingress|egress|nodeport|loadbalancer|external ?ip|\bservices?\b|\bcni\b|\btls\b|mesh|istio|linkerd|dns|target ports?"),
    ("Supply chain & images", r"\bimages?\b|registr|repositor|\btag\b|latest|digest|signature|cosign|notary|sbom|attest|pull ?policy|provenance"),
    ("Workload security (Pod Security)", r"privileg|capabilit|hostpid|hostipc|hostnetwork|host (pid|ipc|network|namespace|path|process)|hostpath|run ?as|root|seccomp|apparmor|selinux|read.?only.*file ?system|securitycontext|security context|sysctl|proc.?mount|volume types?|pod security|psp|docker.?(daemon )?sock|containerd.sock|\bexec\b|shareprocessnamespace|windows|\btty\b|\bssh\b|sandbox|linux hardening|symlink|\\b[ug]id\\b|/etc/hosts"),
    ("Resources, availability & reliability", r"\blimits?\b|requests?\b|replica|probe|liveness|readiness|startup|disruption|\bpdb\b|priority.?class|quota|limitrange|\bhpa\b|topology|affinity|rolling update|ttl|storage ?class|persistentvolume|env var"),
    ("Namespaces & multi-tenancy", r"namespace|tenan|labels?|annotations?"),
    ("Cluster policies & governance", r"polic|admission|webhook|deprecat|api version|crd|sorted keys"),
]
DOMAIN_RES = [(d, re.compile(p, re.I)) for d, p in DOMAINS]

TOPICS: list[tuple[str, str]] = [
    ("Known CVEs / outdated versions", r"cve-\d{4}|outdated"),
    ("File permissions / ownership", r"permission|ownership|chmod|chown"),
    ("Privileged containers", r"privileged(?! ?escalation)|runasprivileged"),
    ("Privilege escalation", r"privilege.?escalation"),
    ("Run as non-root user", r"run(s|ning)? ?as ?(non.?)?root|runasnonroot|runasuser|root user|\buid\b"),
    ("Root group / fsGroup", r"root (primary|group)|runasgroup|fsgroup|supplemental|\bgid\b"),
    ("Host PID namespace", r"host ?pid"),
    ("Host IPC namespace", r"host ?ipc"),
    ("Host network", r"host ?network"),
    ("Host ports", r"host ?port"),
    ("hostPath volumes", r"host ?path"),
    ("Container runtime socket mount", r"docker\.sock|containerd\.sock|crio\.sock|runtime socket|docker socket"),
    ("Linux capabilities", r"capabilit|net_raw|sys_admin|cap_"),
    ("Read-only root filesystem", r"read.?only.?root|readonlyrootfilesystem|read.?only file ?system"),
    ("Seccomp", r"seccomp"),
    ("AppArmor", r"apparmor"),
    ("SELinux", r"selinux"),
    ("Sysctls", r"sysctl"),
    ("procMount", r"proc.?mount"),
    ("Windows HostProcess", r"hostprocess|windows"),
    ("Pod Security Standards / PSA / PSP", r"pod security (standard|admission|polic)|\bpsa\b|\bpsp\b|podsecuritypolicy"),
    ("CPU/memory limits & requests", r"(cpu|memory).{0,20}(limit|request)|(limit|request).{0,20}(cpu|memory)|resource (limit|request|quota)|limitrange"),
    ("Image tag / digest pinning", r"latest|image tag|\btag\b|digest|immutable"),
    ("Image pull policy", r"pull ?policy|alwayspullimages"),
    ("Allowed / trusted registries", r"registr|allowed repos|trusted"),
    ("Image signing & verification", r"signature|signed|cosign|notary|verify.?image|attest|sigstore"),
    ("Liveness / readiness probes", r"probe|liveness|readiness|startup"),
    ("ServiceAccount token automount", r"automount|service ?account token"),
    ("Default ServiceAccount", r"default service ?account"),
    ("Default namespace usage", r"default namespace"),
    ("Secrets in env vars / config", r"secret.{0,30}(env|environment)|env.{0,30}secret|sensitive.{0,20}(env|config)|credentials? in"),
    ("RBAC: access to secrets", r"(access|get|list|watch|read).{0,30}secrets?|secrets?.{0,30}(access|read)"),
    ("RBAC: wildcards", r"wildcard|\*"),
    ("RBAC: cluster-admin", r"cluster.?admin"),
    ("RBAC: exec / attach", r"exec|attach"),
    ("RBAC: create pods / workloads", r"create.{0,20}pods?|pod creation"),
    ("RBAC: impersonate / bind / escalate", r"impersonat|\bbind\b|escalate"),
    ("RBAC: nodes/proxy", r"nodes/proxy|node proxy"),
    ("RBAC: system:masters / anonymous", r"system:masters|system:anonymous|system:unauthenticated|anonymous"),
    ("Network policies", r"network ?polic"),
    ("NodePort / LoadBalancer / external exposure", r"nodeport|loadbalancer|external ?ip|exposed"),
    ("Ingress configuration", r"ingress"),
    ("TLS / certificates / ciphers", r"\btls\b|https|certificate|cert-file|ca-file|cipher|peer-cert|client-cert"),
    ("API / kubelet authn & authz flags", r"authorization-mode|alwaysallow|webhook auth|read-only-port|readonlyport|anonymous-auth|secure-port"),
    ("Kubelet hardening flags", r"streaming-connection|hostname-override|protect-kernel|make-iptables|event-qps|seccomp-default"),
    ("Legacy / risky components (Tiller, dashboard, SSH)", r"tiller|dashboard|\bssh\b|helm v2"),
    ("Profiling endpoints", r"profiling"),
    ("Audit logging", r"audit"),
    ("Admission plugins", r"admission (plugin|control)|noderestriction|eventratelimit|alwaysadmit"),
    ("Encryption at rest / KMS", r"encryption|kms|encrypt"),
    ("etcd TLS / client auth", r"etcd.{0,40}(tls|cert|auth)|peer"),
    ("Certificate rotation", r"rotat"),
    ("Bind address / insecure port", r"bind.?address|insecure.?port|insecure.?bind"),
    ("Token / basic auth files", r"token.?auth.?file|basic.?auth"),
    ("Kernel defaults / event limits", r"protect-kernel-defaults|event.?qps|make-iptables"),
    ("Cloud metadata / IAM", r"metadata|\biam\b|imds|workload identity|irsa"),
    ("Deprecated APIs / versions", r"deprecat|api version|end of life|eol"),
    ("Replicas / PDB / availability", r"replica|disruption|\bpdb\b|hpa|anti.?affinity|topology"),
]
TOPIC_RES = [(t, re.compile(p, re.I)) for t, p in TOPICS]


def classify(r: Rule) -> None:
    hay = " ".join([r.title, r.description, r.benchmark, r.scope, r.rule_id.replace("_", " ")])
    head = " ".join([r.title, r.rule_id.replace("_", " ").replace("-", " ")])
    r.domain = next((d for d, rx in DOMAIN_RES if rx.search(hay)), "General / other")
    # RBAC-kind rules are RBAC regardless of text
    if re.search(r"\b(Cluster)?Role(Binding)?\b", r.resource_kinds) and "Workload" not in r.domain:
        r.domain = "RBAC & identity"
    # prefer title for topics, fall back to full text
    topics = [t for t, rx in TOPIC_RES if rx.search(head)] or [t for t, rx in TOPIC_RES if rx.search(hay)]
    if not r.domain.startswith("RBAC"):
        topics = [t for t in topics if not t.startswith("RBAC:")]
    r.topics = "; ".join(dict.fromkeys(topics[:3]))


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #
EXPORT_FIELDS = ["uid", "source", "rule_id", "title", "severity", "domain", "topics", "check_type", "scope",
                 "resource_kinds", "benchmark", "cis_refs", "frameworks", "description", "rationale",
                 "remediation", "audit", "deprecated_note", "references", "source_url", "raw_severity",
                 *REVIEW_COLUMNS]
HEADERS = {
    "uid": "UID", "source": "Tool", "rule_id": "Rule ID", "title": "Title", "severity": "Severity",
    "domain": "Domain", "topics": "Topic(s)", "check_type": "Check type", "scope": "Scope",
    "resource_kinds": "Resource kinds", "benchmark": "Benchmark / category", "cis_refs": "CIS refs",
    "frameworks": "Frameworks / tags", "description": "Description", "rationale": "Rationale / risk",
    "remediation": "Remediation", "audit": "Audit / check logic", "deprecated_note": "Applicability note",
    "references": "References", "source_url": "Upstream source (pinned)", "raw_severity": "Raw severity",
    "decision": "Decision", "guideline_id": "Company guideline ID", "guideline_text": "Company guideline text",
    "owner": "Owner", "review_notes": "Review notes",
}


def row(r: Rule) -> dict:
    d = asdict(r)
    d["uid"] = r.uid
    return {k: d.get(k, "") for k in EXPORT_FIELDS}


def sort_key(r: Rule):
    return (r.domain, r.topics.split(";")[0] if r.topics else "~", SEVERITY_ORDER.index(r.severity), r.source, r.rule_id)


def write_json(rules: list[Rule], meta: dict, path: Path) -> None:
    path.write_text(json.dumps({"metadata": meta, "rules": [row(r) for r in rules]}, indent=2, ensure_ascii=False))


def write_csv(rules: list[Rule], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=EXPORT_FIELDS)
        w.writerow(HEADERS)
        for r in rules:
            w.writerow(row(r))


def _short(text: str, n: int = 400) -> str:
    text = " ".join(text.split())
    return text if len(text) <= n else text[: n - 1].rsplit(" ", 1)[0] + " …"


def write_markdown(rules: list[Rule], meta: dict, path: Path) -> None:
    out = io.StringIO()
    p = lambda s="": out.write(s + "\n")
    p("# Kubernetes security rules — cross-scanner catalogue")
    p()
    p(f"Generated {meta['generated_at']} from {len(meta['sources'])} scanners, {len(rules)} rules. "
      "Rules are grouped by **domain → topic** so equivalent checks from different tools sit together.")
    p()
    p("| Tool | Rules | Upstream commit |")
    p("|---|---:|---|")
    for s in meta["sources"]:
        p(f"| {s['source']} | {s['rules']} | [`{s['commit'][:10]}`](https://github.com/{s['repo']}/tree/{s['commit']}) |")
    if meta.get("kube_bench_benchmarks"):
        p(f"\nkube-bench benchmarks included: {', '.join(meta['kube_bench_benchmarks'])}")
    by_domain: dict[str, list[Rule]] = defaultdict(list)
    for r in rules:
        by_domain[r.domain].append(r)
    p("\n## Contents\n")
    for d in sorted(by_domain):
        anchor = re.sub(r"[^a-z0-9 -]", "", d.lower()).replace(" ", "-")
        p(f"- [{d}](#{anchor}) — {len(by_domain[d])} rules")
    for d in sorted(by_domain):
        p(f"\n## {d}\n")
        by_topic: dict[str, list[Rule]] = defaultdict(list)
        for r in by_domain[d]:
            by_topic[r.topics.split(";")[0].strip() or "Other"].append(r)
        for t in sorted(by_topic, key=lambda x: (x == "Other", x)):
            group = sorted(by_topic[t], key=lambda r: (SEVERITY_ORDER.index(r.severity), r.source, r.rule_id))
            tools = sorted({r.source for r in group})
            p(f"### {t}  \n_{len(group)} rules · tools: {', '.join(tools)}_\n")
            for r in group:
                bits = [f"`{r.source}` **{r.rule_id}**", r.severity.upper() if r.severity != "unknown" else None,
                        r.check_type if r.check_type == "manual" else None,
                        f"CIS {r.cis_refs}" if r.cis_refs else None]
                p(f"- {' · '.join(b for b in bits if b)} — {r.title}")
                if r.description and r.description != r.title:
                    p(f"  - _What:_ {_short(r.description)}")
                if r.remediation:
                    p(f"  - _Fix:_ {_short(r.remediation, 300)}")
                if r.deprecated_note:
                    p(f"  - ⚠ {r.deprecated_note}")
                p(f"  - [source]({r.source_url})")
            p()
    path.write_text(out.getvalue(), encoding="utf-8")


def write_xlsx(rules: list[Rule], meta: dict, path: Path) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        log.warning("openpyxl not installed — skipping .xlsx (pip install openpyxl)")
        return False

    wb = Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3A5F")
    review_fill = PatternFill("solid", fgColor="FFF4CC")
    sev_fill = {"critical": "C00000", "high": "E26B0A", "medium": "F2C14E", "low": "8DB4E2", "info": "D9D9D9"}
    wrap = Alignment(wrap_text=True, vertical="top")

    # --- README sheet
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Kubernetes security rules — review workbook", True),
        (f"Generated {meta['generated_at']} · {len(rules)} rules from {len(meta['sources'])} scanners", False),
        ("", False),
        ("How to review", True),
        ("1. Work in the 'Rules' sheet. Filter by Domain / Topic to see equivalent checks from different tools side by side.", False),
        ("2. Fill the yellow columns: Decision (Adopt / Adapt / Reject / Not applicable / Needs discussion), Company guideline ID, Company guideline text, Owner, Review notes.", False),
        ("3. Give duplicate rules from different tools the SAME Company guideline ID — the draft step merges them into one guideline and lists every tool rule that verifies it.", False),
        ("4. Use 'Adapt' when the intent applies but the threshold/scope differs for us; write our version in 'Company guideline text'.", False),
        ("5. Run:  python3 k8s_rule_harvester.py draft <this file> -o guideline-draft.md", False),
        ("", False),
        ("Notes", True),
        ("• Severity is normalised from each tool (Kubescape baseScore ≥9 critical, ≥7 high, ≥4 medium; Polaris danger→high, warning→medium). kube-bench (CIS) and Checkov OSS carry no severity → 'unknown'.", False),
        ("• CIS refs are version-specific (e.g. cis-1.12:4.2.1 vs k8s-cis-1.23:4.2.1). Numbering shifts between benchmark versions.", False),
        ("• Domain/Topic are heuristic keyword classifications to help grouping — correct them if wrong.", False),
        ("• Every rule links to the exact upstream file at the fetched commit, so the review is reproducible.", False),
        ("", False),
        ("Sources", True),
    ]
    for text, bold in lines:
        ws.append([text])
        if bold:
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=13 if ws.max_row == 1 else 11)
    ws.append(["Tool", "Repository", "Commit", "Rules", "What it covers"])
    for c in ws[ws.max_row]:
        c.font, c.fill = hdr_font, hdr_fill
    for s in meta["sources"]:
        ws.append([s["source"], f"https://github.com/{s['repo']}", s["commit"], s["rules"], s["description"]])
    if meta.get("kube_bench_benchmarks"):
        ws.append([])
        ws.append(["kube-bench benchmarks", ", ".join(meta["kube_bench_benchmarks"])])
    ws.column_dimensions["A"].width = 26
    for col, w in zip("BCDE", (45, 44, 8, 55)):
        ws.column_dimensions[col].width = w

    # --- Rules sheet
    ws = wb.create_sheet("Rules")
    ws.append([HEADERS[f] for f in EXPORT_FIELDS])
    widths = {"uid": 28, "source": 11, "rule_id": 22, "title": 50, "severity": 10, "domain": 28, "topics": 30,
              "check_type": 10, "scope": 22, "resource_kinds": 22, "benchmark": 20, "cis_refs": 22, "frameworks": 26,
              "description": 60, "rationale": 45, "remediation": 55, "audit": 35, "deprecated_note": 30,
              "references": 35, "source_url": 40, "raw_severity": 14, "decision": 16, "guideline_id": 18,
              "guideline_text": 50, "owner": 14, "review_notes": 40}
    for i, f in enumerate(EXPORT_FIELDS, 1):
        c = ws.cell(1, i)
        c.font, c.alignment = hdr_font, Alignment(wrap_text=True, vertical="center")
        c.fill = PatternFill("solid", fgColor="7F6000") if f in REVIEW_COLUMNS else hdr_fill
        ws.column_dimensions[get_column_letter(i)].width = widths.get(f, 20)
    sev_col = EXPORT_FIELDS.index("severity") + 1
    url_col = EXPORT_FIELDS.index("source_url") + 1
    for r in rules:
        vals = row(r)
        ws.append([str(vals[f])[:32000] for f in EXPORT_FIELDS])
        n = ws.max_row
        for i, f in enumerate(EXPORT_FIELDS, 1):
            cell = ws.cell(n, i)
            cell.alignment = wrap
            if f in REVIEW_COLUMNS:
                cell.fill = review_fill
        if r.severity in sev_fill:
            ws.cell(n, sev_col).fill = PatternFill("solid", fgColor=sev_fill[r.severity])
        ws.cell(n, url_col).hyperlink = r.source_url
        ws.cell(n, url_col).font = Font(color="0563C1", underline="single")
    last = get_column_letter(len(EXPORT_FIELDS))
    tbl = Table(displayName="Rules", ref=f"A1:{last}{ws.max_row}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tbl)
    ws.freeze_panes = "E2"
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True)
    dcol = get_column_letter(EXPORT_FIELDS.index("decision") + 1)
    dv.add(f"{dcol}2:{dcol}{ws.max_row}")
    ws.add_data_validation(dv)
    for n in range(2, ws.max_row + 1):
        ws.row_dimensions[n].height = 60

    # --- Topics sheet: equivalent checks across tools
    ws = wb.create_sheet("Topics x Tools")
    tools = [s["source"] for s in meta["sources"]]
    ws.append(["Topic", "Main domain", "Total", "# tools", *tools])
    grid: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    tdom: dict[str, Counter] = defaultdict(Counter)
    for r in rules:
        t = r.topics.split(";")[0].strip() or f"(no topic) {r.domain}"
        grid[t][r.source].append(r.rule_id)
        tdom[t][r.domain] += 1
    for t, per in sorted(grid.items(), key=lambda kv: (-len(kv[1]), kv[0])):
        ws.append([t, tdom[t].most_common(1)[0][0], sum(len(v) for v in per.values()), len(per),
                   *[", ".join(per.get(tool, []))[:32000] for tool in tools]])
    for c in ws[1]:
        c.font, c.fill = hdr_font, hdr_fill
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 36
    for i in range(5, 5 + len(tools)):
        ws.column_dimensions[get_column_letter(i)].width = 30
    for rw in ws.iter_rows(min_row=2):
        for c in rw:
            c.alignment = wrap
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    # --- CIS cross-reference
    ws = wb.create_sheet("CIS cross-ref")
    ws.append(["CIS control #", "Tool", "Rule ID", "CIS ref (version-specific)", "Title"])
    cis_rows = []
    for r in rules:
        for ref in filter(None, (x.strip() for x in r.cis_refs.split(";"))):
            num = ref.split(":")[-1]
            cis_rows.append((num, r.source, r.rule_id, ref, r.title))
    cis_rows.sort(key=lambda x: ([int(p) if p.isdigit() else 0 for p in x[0].split(".")], x[1]))
    for x in cis_rows:
        ws.append(list(x))
    for c in ws[1]:
        c.font, c.fill = hdr_font, hdr_fill
    for col, w in zip("ABCDE", (14, 12, 24, 30, 80)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # --- Summary
    ws = wb.create_sheet("Summary")
    ws.append(["Domain", "Total", *tools, *SEVERITY_ORDER])
    by_d = defaultdict(list)
    for r in rules:
        by_d[r.domain].append(r)
    for d in sorted(by_d):
        cs, cv = Counter(r.source for r in by_d[d]), Counter(r.severity for r in by_d[d])
        ws.append([d, len(by_d[d]), *[cs.get(t, 0) for t in tools], *[cv.get(s, 0) for s in SEVERITY_ORDER]])
    for c in ws[1]:
        c.font, c.fill = hdr_font, hdr_fill
    ws.column_dimensions["A"].width = 44
    ws.freeze_panes = "B2"

    wb.move_sheet("Summary", offset=-(len(wb.sheetnames) - 2))
    wb.save(path)
    return True


# --------------------------------------------------------------------------- #
# harvest
# --------------------------------------------------------------------------- #
def cmd_harvest(opts) -> int:
    out = Path(opts.output_dir)
    out.mkdir(parents=True, exist_ok=True)
    cache = Path(opts.cache_dir)
    cache.mkdir(parents=True, exist_ok=True)
    refs = dict(x.split("=", 1) for x in (opts.ref or []))
    selected = list(SOURCES) if opts.sources == ["all"] else opts.sources
    unknown = set(selected) - set(SOURCES)
    if unknown:
        log.error("unknown source(s): %s (choose from %s)", ", ".join(unknown), ", ".join(SOURCES))
        return 2

    rules: list[Rule] = []
    meta = {"generated_at": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
            "k8s_version": opts.k8s_version, "sources": [], "errors": []}
    for key in selected:
        spec = SOURCES[key]
        try:
            co = fetch(spec, cache, refs.get(key), opts.refresh, opts.offline)
            found = list(PARSERS[spec.parser](co, opts))
        except Exception as e:  # keep going with other sources
            log.error("%s: %s", key, e)
            meta["errors"].append({"source": key, "error": str(e)})
            continue
        if not found:
            log.warning("%s: 0 rules parsed — upstream layout may have changed (%s)", key, co.path)
        log.info("%-12s %4d rules  @ %s", key, len(found), co.commit[:10])
        meta["sources"].append({"source": key, "repo": spec.repo, "commit": co.commit,
                                "rules": len(found), "description": spec.description})
        rules.extend(found)

    kb = sorted({r.benchmark for r in rules if r.source == "kube-bench"})
    meta["kube_bench_benchmarks"] = kb

    # de-duplicate identical uids (e.g. same kyverno policy in several dirs)
    uniq: dict[str, Rule] = {}
    for r in rules:
        uniq.setdefault(r.uid, r)
    rules = list(uniq.values())
    for r in rules:
        classify(r)
    if opts.min_severity:
        cut = SEVERITY_ORDER.index(opts.min_severity)
        rules = [r for r in rules if r.severity == "unknown" or SEVERITY_ORDER.index(r.severity) <= cut]
    rules.sort(key=sort_key)

    stem = out / opts.basename
    fmts = set(opts.formats)
    written = []
    if "json" in fmts:
        write_json(rules, meta, stem.with_suffix(".json")); written.append(stem.with_suffix(".json"))
    if "csv" in fmts:
        write_csv(rules, stem.with_suffix(".csv")); written.append(stem.with_suffix(".csv"))
    if "md" in fmts:
        write_markdown(rules, meta, stem.with_suffix(".md")); written.append(stem.with_suffix(".md"))
    if "xlsx" in fmts and write_xlsx(rules, meta, stem.with_suffix(".xlsx")):
        written.append(stem.with_suffix(".xlsx"))

    print(f"\n{len(rules)} rules from {len(meta['sources'])} sources")
    for s in meta["sources"]:
        print(f"  {s['source']:<12} {s['rules']:>5}")
    for e in meta["errors"]:
        print(f"  !! {e['source']}: {e['error']}")
    print("Written:\n  " + "\n  ".join(map(str, written)))
    return 1 if meta["errors"] and not rules else 0


# --------------------------------------------------------------------------- #
# draft — reviewed workbook -> company guideline (Markdown)
# --------------------------------------------------------------------------- #
def _read_reviewed(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8") as fh:
            rd = csv.reader(fh)
            header = next(rd)
            inv = {v: k for k, v in HEADERS.items()}
            keys = [inv.get(h, h) for h in header]
            return [dict(zip(keys, r)) for r in rd]
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True)["Rules"]
    rows = ws.iter_rows(values_only=True)
    inv = {v: k for k, v in HEADERS.items()}
    keys = [inv.get(h, h) for h in next(rows)]
    return [{k: ("" if v is None else str(v)) for k, v in zip(keys, r)} for r in rows]


def cmd_draft(opts) -> int:
    rows = _read_reviewed(Path(opts.workbook))
    keep = [r for r in rows if r.get("decision", "").strip().lower() in ("adopt", "adapt")]
    if not keep:
        print("No rows with Decision = Adopt/Adapt found. Review the 'Rules' sheet first.")
        return 1
    abbrev = {}
    auto: dict[str, int] = defaultdict(int)
    groups: dict[str, list[dict]] = defaultdict(list)
    for r in keep:
        gid = r.get("guideline_id", "").strip()
        if not gid:  # auto id from domain initials
            dom = r.get("domain", "GEN")
            ab = abbrev.setdefault(dom, "".join(w[0] for w in re.findall(r"[A-Za-z]+", dom))[:4].upper() or "GEN")
            auto[ab] += 1
            gid = f"{opts.prefix}-{ab}-{auto[ab]:03d}"
        groups[gid].append(r)

    by_domain: dict[str, list[str]] = defaultdict(list)
    for gid, rs in groups.items():
        by_domain[Counter(r.get("domain", "") for r in rs).most_common(1)[0][0]].append(gid)

    sev_rank = {s: i for i, s in enumerate(SEVERITY_ORDER)}
    out = io.StringIO()
    p = lambda s="": out.write(s + "\n")
    p(f"# {opts.title}")
    p()
    p(f"_Draft generated {dt.date.today().isoformat()} from {len(keep)} reviewed scanner rules → {len(groups)} guidelines._")
    p()
    p("**Requirement levels:** MUST = enforced / blocking, SHOULD = expected unless a documented exception exists.")
    p()
    for d in sorted(by_domain):
        p(f"## {d}\n")
        for gid in sorted(by_domain[d]):
            rs = sorted(groups[gid], key=lambda r: sev_rank.get(r.get("severity", "unknown"), 9))
            text = next((r["guideline_text"] for r in rs if r.get("guideline_text", "").strip()), rs[0].get("title", ""))
            sev = rs[0].get("severity", "unknown")
            level = "MUST" if sev in ("critical", "high") else "SHOULD"
            owner = ", ".join(sorted({r["owner"] for r in rs if r.get("owner")}))
            p(f"### {gid} — {text}\n")
            p(f"- **Level:** {level}  ·  **Severity:** {sev}" + (f"  ·  **Owner:** {owner}" if owner else ""))
            rat = next((r.get("rationale") or r.get("description") for r in rs if (r.get("rationale") or r.get("description"))), "")
            if rat:
                p(f"- **Why:** {_short(rat, 600)}")
            rem = next((r["remediation"] for r in rs if r.get("remediation")), "")
            if rem:
                p(f"- **How to comply:**\n\n  ```\n  " + "\n  ".join(rem.strip().splitlines()[:25]) + "\n  ```")
            cis = sorted({c for r in rs for c in r.get("cis_refs", "").split(";") if c.strip()})
            if cis:
                p(f"- **CIS:** {', '.join(c.strip() for c in cis)}")
            p("- **Verified by:** " + "; ".join(f"{r['source']} `{r['rule_id']}`" for r in rs))
            notes = [r["review_notes"] for r in rs if r.get("review_notes")]
            if notes:
                p(f"- **Review notes:** {' / '.join(notes)}")
            p()
    Path(opts.output).write_text(out.getvalue(), encoding="utf-8")
    print(f"{len(groups)} guidelines written to {opts.output}")
    return 0


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-v", "--verbose", action="store_true")
    sub = ap.add_subparsers(dest="cmd")

    h = sub.add_parser("harvest", help="fetch, normalise and export rules (default)")
    h.add_argument("-s", "--sources", nargs="+", default=["all"], help=f"sources: all | {' '.join(SOURCES)}")
    h.add_argument("-o", "--output-dir", default="k8s-rules-out")
    h.add_argument("--basename", default="k8s-security-rules")
    h.add_argument("--cache-dir", default=".k8s-rule-cache", help="where upstream repos are cloned")
    h.add_argument("--formats", nargs="+", default=["xlsx", "md", "csv", "json"], choices=["xlsx", "md", "csv", "json"])
    h.add_argument("--kube-bench-benchmarks", nargs="+", default=["latest"],
                   help="latest | all | explicit names e.g. cis-1.12 eks-1.8.0 aks-1.8 gke-1.9.0 rke2-cis-1.9")
    h.add_argument("--k8s-version", help="e.g. 1.30 — picks the matching kube-bench CIS benchmark via its version_mapping")
    h.add_argument("--ref", action="append", metavar="SOURCE=REF",
                   help="pin a source to a tag/branch, e.g. --ref kube-bench=v0.10.6 (repeatable)")
    h.add_argument("--min-severity", choices=SEVERITY_ORDER[:-1],
                   help="drop rules below this severity (rules with unknown severity are always kept)")
    h.add_argument("--refresh", action="store_true", help="re-clone sources even if cached")
    h.add_argument("--offline", action="store_true", help="use cached clones only")

    d = sub.add_parser("draft", help="generate a draft guideline from a reviewed workbook/CSV")
    d.add_argument("workbook", help="reviewed .xlsx (Rules sheet) or .csv")
    d.add_argument("-o", "--output", default="k8s-security-guideline-draft.md")
    d.add_argument("--prefix", default="K8S", help="guideline ID prefix for rows without an ID")
    d.add_argument("--title", default="Kubernetes Cluster Security Guideline (DRAFT)")

    sub.add_parser("list-sources", help="show supported scanners")

    argv = sys.argv[1:] if argv is None else argv
    if not argv or argv[0].startswith("-") and argv[0] not in ("-h", "--help", "-v", "--verbose"):
        argv = ["harvest", *argv]
    opts = ap.parse_args(argv)
    if opts.cmd is None:
        opts = ap.parse_args([*argv, "harvest"])
    logging.basicConfig(level=logging.DEBUG if opts.verbose else logging.INFO, format="%(levelname)s %(message)s")

    if opts.cmd == "list-sources":
        for s in SOURCES.values():
            print(f"{s.key:<12} github.com/{s.repo:<40} {s.description}")
        return 0
    if opts.cmd == "draft":
        return cmd_draft(opts)
    return cmd_harvest(opts)


if __name__ == "__main__":
    sys.exit(main())